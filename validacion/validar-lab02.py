r"""Valida el laboratorio 02: modelo, clasificador y el archivo .xlsx exportado."""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

HTML = Path(__file__).resolve().parent.parent / "lab02-resorte-rlc.html"
TMP = Path(__file__).resolve().parent / "_tmp"
NODE = shutil.which("node") or sys.exit("No se encontró node en el PATH")
XLSX = TMP / "exportado.xlsx"

script = max(re.findall(r"<script>(.*?)</script>", HTML.read_text(encoding="utf-8"), re.S), key=len)
TMP.mkdir(exist_ok=True)
(TMP / "completo2.js").write_text(script, encoding="utf-8")

print("--- sintaxis (node --check) ---")
r = subprocess.run([NODE, "--check", str(TMP / "completo2.js")], capture_output=True, text=True)
print((r.stdout + r.stderr).strip() or "OK")
if r.returncode != 0:
    sys.exit("ERROR DE SINTAXIS")

math = script[script.index("/* ==== MATH:BEGIN ==== */"):script.index("/* ==== MATH:END ==== */")]
export = script[script.index("const MIME_XLSX"):script.index("function nombreArchivo")]

pruebas = r"""
const fs = require('fs');
let fallos = 0;
const check = (n, c, extra = "") => {
  console.log((c ? "  ok    " : "  FALLA ") + n + (extra ? "  " + extra : ""));
  if (!c) fallos++;
};
const cerca = (a, b, tol) => Math.abs(a - b) <= tol;
const base = { m: 1, c: 0.5, k: 10, F0: 2, w: 2, x0: 1, v0: 0, efecto: 'ninguno',
               Fc: 0.5, ablanda: 0.35, c2rel: 0.8 };

// 1. Clasificación
check("zeta<1 subamortiguado", clasificar(1, 0.5, 10).regimen === 'subamortiguado');
check("zeta=1 crítico", clasificar(1, 2 * Math.sqrt(10), 10).regimen === 'crítico');
check("zeta>1 sobreamortiguado", clasificar(1, 8, 10).regimen === 'sobreamortiguado');
check("omega0 y delta coherentes",
      cerca(clasificar(1, 0.5, 10).w0, Math.sqrt(10), 1e-12) &&
      clasificar(1, 0.5, 10).delta < 0 && clasificar(1, 8, 10).delta > 0);

// 2. Amplitud permanente teoría vs integración
let peorAmp = 0;
for (const w of [1.0, 1.8, 2.6, 3.2]) {
  const Q = Object.assign({}, base, { w: w });
  const teoria = amplitudEstacionaria(Q.m, Q.c, Q.k, Q.F0, w).A;
  const pts = integrarEstado(Q, 'ninguno', 0, [0, 0], 120, 0.005);
  let num = 0;
  for (let i = Math.floor(pts.length * 0.75); i < pts.length; i++) num = Math.max(num, Math.abs(pts[i][1]));
  peorAmp = Math.max(peorAmp, Math.abs(num - teoria) / teoria);
}
check("amplitud teórica = amplitud integrada (<2%)", peorAmp < 0.02, `peor=${(peorAmp * 100).toFixed(2)}%`);

// 3. Resonancia
const wr = frecuenciaResonancia(base.m, base.c, base.k);
const barrido = barridoResonancia(Object.assign({}, base, { efecto: 'ninguno' }), 40);
const pico = barrido.reduce((a, b) => (b.A > a.A ? b : a));
check("pico numérico en omega_r (<3%)", Math.abs(pico.w - wr) / wr < 0.03,
      `pico=${pico.w.toFixed(3)} omega_r=${wr.toFixed(3)}`);
check("pico vale A_max teórica (<5%)",
      Math.abs(pico.A - amplitudEstacionaria(base.m, base.c, base.k, base.F0, wr).A) /
      amplitudEstacionaria(base.m, base.c, base.k, base.F0, wr).A < 0.05);
check("con zeta alto no hay pico de resonancia", frecuenciaResonancia(1, 5, 10) === null,
      `zeta=${clasificar(1, 5, 10).zeta.toFixed(3)}`);

// 4. Solución analítica libre vs RK4
let peorLibre = 0;
for (const [m, c, k] of [[1, 0.5, 10], [1, 2 * Math.sqrt(10), 10], [1, 8, 10]]) {
  const pts = integrarEstado({ m, c, k, F0: 0, w: 0, x0: 1, v0: 0.5 }, 'ninguno', 0, [1, 0.5], 3, 0.001);
  for (const [t, x] of pts) peorLibre = Math.max(peorLibre, Math.abs(x - analiticaLibre(m, c, k, 1, 0.5, t)));
}
check("solución analítica = RK4 (<1e-4)", peorLibre < 1e-4, `peor=${peorLibre.toExponential(2)}`);

// 5. Energía y corriente
const libre = integrarEstado({ m: 1, c: 0, k: 10, F0: 0, w: 0, x0: 1, v0: 0 }, 'ninguno', 0, [1, 0], 20, 0.001);
const E = (x, v) => 0.5 * v * v + 0.5 * 10 * x * x;
const deriva = libre.reduce((m, p) => Math.max(m, Math.abs(E(p[1], p[2]) - E(1, 0))), 0);
check("energía conservada sin amortiguamiento (<1e-6)", deriva < 1e-6);
const conQ = integrarEstado(Object.assign({}, base, { x0: 0, v0: 0 }), 'ninguno', 0, [0, 0], 60, 0.005);
let aq = 0, ai = 0;
for (let i = Math.floor(conQ.length * 0.8); i < conQ.length; i++) {
  aq = Math.max(aq, Math.abs(conQ[i][1])); ai = Math.max(ai, Math.abs(conQ[i][2]));
}
check("la corriente i = q' va a omega por la carga (<3%)", Math.abs(ai - base.w * aq) / (base.w * aq) < 0.03);

// 6. Clasificador de efectos (la verificación)
for (const [efecto, que] of [['ninguno', 'modelo lineal'], ['seco', 'rozamiento seco'],
                             ['no_lineal', 'resorte no lineal'], ['cuadratica', 'amort. cuadrático']]) {
  const d = diagnosticar(Object.assign({}, base, { efecto: efecto }));
  check(`diagnostica «${que}»`, d.efecto === efecto,
        `dijo «${d.efecto}» · ${d.evidencia.slice(0, 58)}…`);
}

// 6b. Clasificador en los casos que antes fallaban: sin excitación y sin pico de resonancia
for (const efecto of EFECTOS) {
  const d = diagnosticar(Object.assign({}, base, { F0: 0, efecto: efecto }));
  check(`con F0=0 diagnostica «${efecto}»`, d.efecto === efecto, `dijo «${d.efecto}»`);
}
const cCrit = 2 * Math.sqrt(base.k * base.m);
for (const [nombre, c] of [['crítico', cCrit], ['sobreamortiguado', 2.4 * cCrit], ['zeta=0.75 (sin pico)', 1.5 * Math.sqrt(base.k * base.m)]]) {
  const d = diagnosticar(Object.assign({}, base, { c: c, efecto: 'ninguno' }));
  check(`régimen ${nombre} sin efecto: no inventa uno`, d.efecto === 'ninguno', `dijo «${d.efecto}»`);
}

// 6c. Barrido de robustez: aciertos y, sobre todo, ningún efecto inventado ni confundido
{
  let total = 0, aciertos = 0, inventados = 0, confundidos = 0;
  for (const m of [0.5, 1, 3]) for (const k of [2, 10, 30]) for (const zeta of [0.03, 0.1, 0.2])
  for (const F0 of [0, 2]) for (const x0 of [0.3, -1.5]) for (const efecto of EFECTOS) {
    const P = Object.assign({}, base, { m, k, c: 2 * zeta * Math.sqrt(k * m), F0, w: 0.9 * Math.sqrt(k / m), x0, efecto });
    const d = diagnosticar(P).efecto;
    total++;
    if (d === efecto) aciertos++;
    else if (efecto === 'ninguno') inventados++;
    else if (d !== 'ninguno') confundidos++;
  }
  check("barrido del clasificador: al menos 95 % de aciertos", aciertos / total >= 0.95, `${aciertos}/${total}`);
  check("barrido del clasificador: nunca inventa un efecto en el modelo lineal", inventados === 0, `${inventados}`);
  check("barrido del clasificador: nunca confunde un efecto con otro", confundidos === 0, `${confundidos}`);
}

// 6d. El resorte que se ablanda no diverge (antes llegaba a 1e38 en menos de 1 s)
for (const [x0, v0, F0, w, c] of [[1.5, 3, 0, 2, 0], [0, 0, 4, 2.6, 0.3], [2, 3, 10, 1, 0]]) {
  const Q = Object.assign({}, base, { F0, w, c, efecto: 'no_lineal' });
  const pts = integrarEstado(Q, 'no_lineal', 0, [x0, v0], 60, 1 / 240);
  const fin = pts[pts.length - 1];
  check(`resorte no lineal acotado (x0=${x0}, v0=${v0}, F0=${F0})`,
        cerca(fin[0], 60, 1e-6) && pts.every(p => Number.isFinite(p[1]) && Math.abs(p[1]) < 1e3),
        `t final=${fin[0].toFixed(2)}`);
}
{
  const Q = Object.assign({}, base, { c: 0, F0: 0, efecto: 'no_lineal' });
  const pts = integrarEstado(Q, 'no_lineal', 0, [1.5, 0], 20, 0.001);
  const En = p => 0.5 * Q.m * p[2] * p[2] + energiaPotencial(Q, 'no_lineal', p[1]);
  const E0 = En(pts[0]);
  const deriva = pts.reduce((m, p) => Math.max(m, Math.abs(En(p) - E0)), 0);
  check("energía del resorte no lineal conservada sin amortiguamiento (<1e-6)", deriva < 1e-6,
        `deriva=${deriva.toExponential(2)}`);
}

// 6e. Rozamiento seco: la masa se queda pegada de verdad (sin vibrar) dentro de la banda |k x| <= Fc
{
  const Q = Object.assign({}, base, { F0: 0, efecto: 'seco' });
  const pts = integrarEstado(Q, 'seco', 0, [1, 0], 30, 1 / 240);
  const cola = pts.slice(Math.floor(pts.length * 0.8));
  const fin = pts[pts.length - 1];
  const recorrido = Math.max(...cola.map(p => p[1])) - Math.min(...cola.map(p => p[1]));
  check("rozamiento seco: se detiene sin vibrar", recorrido < 1e-6, `recorrido=${recorrido.toExponential(2)}`);
  check("rozamiento seco: se detiene donde el resorte no vence a Fc", Q.k * Math.abs(fin[1]) <= Q.Fc + 1e-9,
        `x=${fin[1].toFixed(4)}`);
}

// 7. Exportación a Excel: se genera el archivo y se revisa
(async () => {
  const filas = (n, f) => Array.from({ length: n }, (_, i) => f(i));
const hojas = [
  { nombre: 'Datos', filas: [['t (s)', 'x fenomeno (m)', 'x modelo (m)', 'diferencia (fen - mod)'],
                             [0, 1, 0.999, 0.001], [0.5, 0.7, 0.699, 0.001], [1, 0.2, 0.199, 0.001]] },
  { nombre: 'Parametros', filas: [['parametro', 'valor', 'unidad'], ['masa', 1, 'kg'], ['zeta', 0.079, '']] },
  { nombre: 'Metricas', filas: [['metrica', 'valor'], ['discrepancia', 0.42], ['RMS', 0.0013]] },
  { nombre: 'Resonancia', filas: [['omega', 'medida', 'modelo'],
                                  ...filas(20, i => [0.2 + i * 0.1, 1 / (1 + i), 1.2 / (1 + i * 0.5)])] }
];
const blob = construirXLSX(hojas);
const vista = await blob.arrayBuffer();
fs.writeFileSync(ARCHIVO_XLSX, Buffer.from(vista));
check("el .xlsx se genera con contenido", vista.byteLength > 1500, `${vista.byteLength} bytes`);
check("firma ZIP correcta (PK\\x03\\x04)",
      Buffer.from(vista).slice(0, 4).toString('hex') === '504b0304');
check("contiene [Content_Types].xml",
      Buffer.from(vista).includes(Buffer.from('[Content_Types].xml')));
check("contiene las cuatro hojas",
      ['sheet1.xml', 'sheet2.xml', 'sheet3.xml', 'sheet4.xml'].every(s => Buffer.from(vista).includes(Buffer.from(s))));

console.log(fallos === 0 ? "\nTODO OK" : `\n${fallos} FALLA(S)`);
process.exit(fallos === 0 ? 0 : 1);
})();
"""

(TMP / "pruebas2.js").write_text(
    f"const ARCHIVO_XLSX = {str(XLSX)!r};\n" + math + export + pruebas, encoding="utf-8")
print("\n--- aserciones ---")
r2 = subprocess.run([NODE, "--no-warnings", str(TMP / "pruebas2.js")], capture_output=True, text=True)
print(r2.stdout.strip())
if r2.returncode != 0:
    print("stderr:", r2.stderr.strip()[:900])
    sys.exit(r2.returncode)

# --- Verificación con openpyxl del archivo realmente exportado ---
print("\n--- el .xlsx abierto con openpyxl ---")
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(XLSX, data_only=True)
print("  hojas:", wb.sheetnames)
ok = True
def revisar(nombre, cond, extra=""):
    global ok
    print(("  ok    " if cond else "  FALLA ") + nombre + ("  " + extra if extra else ""))
    ok = ok and cond

revisar("cuatro hojas con los nombres previstos",
        wb.sheetnames == ['Datos', 'Parametros', 'Metricas', 'Resonancia'], str(wb.sheetnames))
ws = wb['Datos']
revisar("encabezado de la hoja Datos", ws['A1'].value == 't (s)' and ws['C1'].value == 'x modelo (m)')
revisar("los números son numéricos, no texto",
        isinstance(ws['A2'].value, (int, float)) and isinstance(ws['B2'].value, (int, float)),
        f"A2={ws['A2'].value!r} ({type(ws['A2'].value).__name__})")
revisar("las filas de datos están completas", ws.max_row == 4 and ws.max_column == 4,
        f"{ws.max_row}×{ws.max_column}")
wsr = wb['Resonancia']
revisar("la hoja Resonancia trae el barrido", wsr.max_row == 21, f"{wsr.max_row} filas")
revisar("celdas vacías se leen como None", wb['Metricas']['C3'].value is None)

print("\nRESULTADO:", "exportación verificada" if ok else "revisar la exportación")

# --- Revisión estática de la página: errores de contenido que ya aparecieron una vez ---
print("\n--- contenido de la página ---")
texto = HTML.read_text(encoding="utf-8")
cmax = float(re.search(r'id="r2"[^>]*max="([\d.]+)"', texto).group(1))
revisar("el deslizador de c alcanza el amortiguamiento crítico de todo el rango",
        cmax >= 2 * (5 * 40) ** 0.5, f"max={cmax} frente a 2·sqrt(5·40)={2 * (5 * 40) ** 0.5:.2f}")
revisar("el RLC no usa mC ni mA (las ecuaciones están en C y A)",
        not re.search(r"\bm[CA]\b", texto))
revisar("ln(A_n) se grafica contra t_n, no contra el número de ciclo",
        "contra el número de ciclo" not in texto and "contra el instante" in texto)
revisar("el modo examen no escribe el efecto en el selector al empezar",
        'value="oculto"' in texto and "selEfecto.value = 'oculto'" in texto)
revisar("la historia guarda 20 s (4800 pasos de 1/240 s)", "HIST_MAX = 4800" in texto)

print("\nRESULTADO:", "laboratorio verificado" if ok else "revisar el laboratorio")
sys.exit(0 if ok else 1)


